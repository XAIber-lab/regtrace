import argparse
import copy
import json
import re
import subprocess
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from contextualization_confidence import (
    CONFIDENCE_LEVELS,
    CONFIDENCE_WEIGHTS,
    SELF_CONSISTENCY_STATUSES,
    SelfConsistencyEvidence,
    build_self_consistency_evidence,
    clean_action_conflicts_with_review,
    change_confidence_summary,
    clean_action_is_publishable,
    confidence_display,
    review_needed_target_keys,
    score_action_confidence,
)
from project_paths import PROMPTS_DIR, SOURCE_DIR

DEFAULT_PROMPT = PROMPTS_DIR / "update_contextualization_prompt.txt"
DEFAULT_QUERY_SCRIPT = SOURCE_DIR / "query.py"
DEFAULT_OUTPUT_LANGUAGE = "it"

SUBCATEGORY_RE = re.compile(r"\b(?:DP-)?[A-Z]{2}\.[A-Z]{2}-\d{1,2}\b", re.IGNORECASE)
SUBCATEGORY_ID_RE = re.compile(
    r"(?P<family>(?:DP-)?[A-Z]{2}\.[A-Z]{2})-(?P<number>\d{1,2})",
    re.IGNORECASE,
)
CIRC285_RE = re.compile(
    r"(?:circ(?:olare)?\.?\s*(?:n(?:o|\.|\W)?\s*)?285\b|circ\s*285\b|\b285\b)",
    re.IGNORECASE,
)

ACTION_FILL_COLORS = {
    "new": "C6EFCE",
    "modify": "FFEB9C",
    "deprecate": "F4CCCC",
    "review_needed": "D9EAD3",
}
ACTION_LABELS = {
    "new": "Nuovo",
    "modify": "Modifica",
    "deprecate": "Deprecazione",
    "no_change": "Nessuna modifica",
    "review_needed": "Da verificare",
}
RELEVANCE_LABELS = {
    "relevant": "Rilevante",
    "not_relevant": "Non rilevante",
    "uncertain": "Incerto",
}
MERGE_STAT_KEYS = {
    "new": "new_rows",
    "modify": "modified_rows",
    "deprecate": "deprecated_rows",
    "review_needed": "review_needed_rows",
}
MERGE_STAT_LABELS = [
    ("Nuove righe", "new_rows"),
    ("Righe modificate", "modified_rows"),
    ("Righe deprecate", "deprecated_rows"),
    ("Righe da verificare", "review_needed_rows"),
]
APPENDABLE_ACTIONS = {"new", "modify", "review_needed"}


@dataclass(frozen=True)
class RowResolution:
    row_idx: int | None
    review_status: str
    requires_review: bool = False


@dataclass(frozen=True)
class MergeReviewEvent:
    change_id: str
    workbook: str
    original_action: str
    effective_action: str
    row_idx: int | None
    subject: str
    review_status: str


def language_display_name(output_language: str) -> str:
    return {"it": "Italian", "en": "English"}.get(output_language.lower(), output_language)


def contextualization_final_reminder(output_language: str) -> str:
    language_name = language_display_name(output_language)
    return f"""

FINAL CHECK BEFORE ANSWERING:
Return valid JSON only, conforming to the schema.
Write newly generated workbook-facing content in {language_name}.
Copy canonical descriptions and preserved control wording verbatim, even in another language.
Set the top-level language field to exactly {output_language!r}.
Use only subcategory codes present in the framework core input.
Do not map changes that are not cybersecurity-related: classify them as not_relevant and leave actions empty.
General governance, risk management, controls, models, data, reporting, or documentation
language is not cybersecurity by itself; require an explicit ICT/security/privacy or
digital operational resilience connection.
Use prototype_actions.existing_row only when it is copied from contextualization_prototype_rows
and that row has the same subcategory_id; never use framework_core.row_number as existing_row.
Use control_actions.existing_control_row only when it is copied from existing_control_rows
and its current subcategories overlap the proposed subcategories_with_weight.
For reference-only changes, prefer no_change or review_needed unless an operational ICT requirement changes.
Use review_needed for unresolved applicability or insufficient evidence only when a plausible target
is supported by the supplied rows; otherwise leave the corresponding action array empty.
For substantive actions, identify the changed requirement and justify the target and action in the
rationale, pointing only to supporting locations or passages present in the supplied evidence.
Use deprecate only when the annotated change explicitly removes, supersedes, or invalidates
an existing obligation/control; never deprecate only because a control is generic.
"""


def schema(
    allowed_subcategory_ids: list[str] | None = None,
    expected_change_ids: list[str] | None = None,
    output_language: str = DEFAULT_OUTPUT_LANGUAGE,
    include_confidence_fields: bool = False,
) -> dict[str, Any]:
    action_enum = ["new", "modify", "deprecate", "no_change", "review_needed"]
    class_enum = ["Obbligatoria", "Consigliata", "Facoltativa", "Non applicabile", "Da verificare"]
    priority_enum = ["ALTA", "MEDIA", "BASSA", "Non applicabile", "Da verificare"]
    subcategory_id_schema: dict[str, Any] = {"type": "string"}
    change_id_schema: dict[str, Any] = {"type": "string"}
    confidence_component_schema = {
        "type": "object",
        "additionalProperties": False,
        "required": ["criterion", "earned_points", "max_points", "status"],
        "properties": {
            "criterion": {"type": "string"},
            "earned_points": {"type": "integer", "minimum": 0},
            "max_points": {"type": "integer", "minimum": 0},
            "status": {"type": "string", "enum": ["pass", "partial", "fail"]},
        },
    }
    confidence_assessment_schema = {
        "type": "object",
        "additionalProperties": False,
        "required": [
            "score",
            "level",
            "deterministic_score",
            "self_consistency_bonus",
            "self_consistency_status",
            "components",
        ],
        "properties": {
            "score": {"type": "integer", "minimum": 0, "maximum": 100},
            "level": {"type": "string", "enum": sorted(CONFIDENCE_LEVELS)},
            "deterministic_score": {"type": "integer", "minimum": 0, "maximum": 90},
            "self_consistency_bonus": {"type": "integer", "minimum": 0, "maximum": 10},
            "self_consistency_status": {
                "type": "string",
                "enum": sorted(SELF_CONSISTENCY_STATUSES),
            },
            "components": {"type": "array", "items": confidence_component_schema},
        },
    }
    confidence_action_property = (
        {"confidence_assessment": confidence_assessment_schema}
        if include_confidence_fields
        else {}
    )
    confidence_top_property = (
        {
            "confidence_scoring": {
                "type": "object",
                "additionalProperties": False,
                "required": ["method", "version", "self_consistency_requested", "weights"],
                "properties": {
                    "method": {"type": "string"},
                    "version": {"type": "string"},
                    "self_consistency_requested": {"type": "boolean"},
                    "weights": {
                        "type": "object",
                        "additionalProperties": {"type": "integer", "minimum": 0},
                    },
                },
            }
        }
        if include_confidence_fields
        else {}
    )

    if allowed_subcategory_ids:
        subcategory_id_schema["enum"] = allowed_subcategory_ids
    if expected_change_ids:
        change_id_schema["enum"] = expected_change_ids

    return {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "name": "contextualization_update",
        "type": "object",
        "additionalProperties": False,
        "required": [
            "contextualization_id",
            "document_id",
            "source_annotation_id",
            "language",
            "change_contextualizations",
            "merge_summary",
        ],
        "properties": {
            "contextualization_id": {"type": "string"},
            "document_id": {"type": "string"},
            "source_annotation_id": {"type": "string"},
            "language": {"type": "string", "enum": [output_language]},
            "change_contextualizations": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": [
                        "change_id",
                        "cybersecurity_relevance",
                        "relevance_rationale",
                        "classe",
                        "priority",
                        "framework_mappings",
                        "prototype_actions",
                        "control_actions",
                        "notes",
                    ],
                    "properties": {
                        "change_id": change_id_schema,
                        "cybersecurity_relevance": {
                            "type": "string",
                            "enum": ["relevant", "not_relevant", "uncertain"],
                        },
                        "relevance_rationale": {"type": "string"},
                        "classe": {"type": "string", "enum": class_enum},
                        "priority": {"type": "string", "enum": priority_enum},
                        "framework_mappings": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "required": [
                                    "subcategory_id",
                                    "subcategory_text",
                                    "mapping_weight",
                                    "rationale",
                                ],
                                "properties": {
                                    "subcategory_id": subcategory_id_schema,
                                    "subcategory_text": {"type": "string"},
                                    "mapping_weight": {
                                        "type": "string",
                                        "enum": ["ALTO", "MEDIO", "BASSO"],
                                    },
                                    "rationale": {"type": "string"},
                                },
                            },
                        },
                        "prototype_actions": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "required": [
                                    "action",
                                    "subcategory_id",
                                    "subcategory_text",
                                    "existing_row",
                                    "classe",
                                    "priority",
                                    "source_reference",
                                    "rationale",
                                ],
                                "properties": {
                                    "action": {"type": "string", "enum": action_enum},
                                    "subcategory_id": subcategory_id_schema,
                                    "subcategory_text": {"type": "string"},
                                    "existing_row": {"type": ["integer", "null"]},
                                    "classe": {"type": "string", "enum": class_enum},
                                    "priority": {"type": "string", "enum": priority_enum},
                                    "source_reference": {"type": ["string", "null"]},
                                    "rationale": {"type": "string"},
                                    **confidence_action_property,
                                },
                            },
                        },
                        "control_actions": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "required": [
                                    "action",
                                    "existing_control_row",
                                    "existing_control_id",
                                    "control_text",
                                    "subcategories_with_weight",
                                    "references",
                                    "priority",
                                    "delegated_to",
                                    "rationale",
                                ],
                                "properties": {
                                    "action": {"type": "string", "enum": action_enum},
                                    "existing_control_row": {"type": ["integer", "null"]},
                                    "existing_control_id": {"type": ["string", "null"]},
                                    "control_text": {"type": ["string", "null"]},
                                    "subcategories_with_weight": {"type": ["string", "null"]},
                                    "references": {"type": ["string", "null"]},
                                    "priority": {"type": "string", "enum": priority_enum},
                                    "delegated_to": {"type": ["string", "null"]},
                                    "rationale": {"type": "string"},
                                    **confidence_action_property,
                                },
                            },
                        },
                        "notes": {"type": ["string", "null"]},
                    },
                },
            },
            "merge_summary": {
                "type": "object",
                "additionalProperties": False,
                "required": [
                    "relevant_changes",
                    "non_relevant_changes",
                    "uncertain_changes",
                    "new_items",
                    "modifications",
                    "deprecations",
                    "summary",
                ],
                "properties": {
                    "relevant_changes": {"type": "integer"},
                    "non_relevant_changes": {"type": "integer"},
                    "uncertain_changes": {"type": "integer"},
                    "new_items": {"type": "integer"},
                    "modifications": {"type": "integer"},
                    "deprecations": {"type": "integer"},
                    "summary": {"type": "string"},
                },
            },
            **confidence_top_property,
        },
    }


def value_to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def compact_text(value: Any, limit: int = 1600) -> str:
    text = re.sub(r"\s+", " ", value_to_text(value))
    if not limit or len(text) <= limit:
        return text
    return text[: limit - 15].rstrip() + " ...[truncated]"


def normalize(value: Any) -> str:
    text = value_to_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_subcategory_id(value: Any) -> str:
    text = value_to_text(value)
    match = SUBCATEGORY_RE.search(text)
    if match:
        return canonicalize_subcategory_id(match.group(0))
    if ":" in text:
        return canonicalize_subcategory_id(text.split(":", 1)[0].strip())
    return canonicalize_subcategory_id(text.strip())


def canonicalize_subcategory_id(value: Any) -> str:
    text = value_to_text(value).strip()
    match = SUBCATEGORY_ID_RE.fullmatch(text)
    if not match:
        return text
    return f"{match.group('family').upper()}-{int(match.group('number')):02d}"


def subcategory_id_aliases(value: Any) -> set[str]:
    text = value_to_text(value)
    match = SUBCATEGORY_RE.search(text)
    if match:
        original = match.group(0)
    elif ":" in text:
        original = text.split(":", 1)[0].strip()
    else:
        original = text.strip()

    canonical = canonicalize_subcategory_id(original)
    return {subcategory_id for subcategory_id in [original, canonical] if subcategory_id}


def subcategory_id_aliases_in_text(value: Any) -> set[str]:
    aliases: set[str] = set()
    for match in SUBCATEGORY_RE.finditer(value_to_text(value)):
        aliases.update(subcategory_id_aliases(match.group(0)))
    return aliases


def normalize_subcategory_references(value: Any) -> str:
    return SUBCATEGORY_RE.sub(
        lambda match: canonicalize_subcategory_id(match.group(0)),
        value_to_text(value),
    )


def parse_column_ref(ref: str | None) -> int | None:
    if not ref:
        return None
    cleaned = ref.strip()
    if not cleaned:
        return None
    if cleaned.isdigit():
        return int(cleaned)
    return column_index_from_string(cleaned.upper())


def find_header_row(ws: Any, required_terms: list[str], max_scan_rows: int = 12) -> int:
    normalized_terms = [normalize(term) for term in required_terms]
    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        row_text = " ".join(normalize(ws.cell(row_idx, col_idx).value) for col_idx in range(1, ws.max_column + 1))
        if all(term in row_text for term in normalized_terms):
            return row_idx
    return 1


def header_map(ws: Any, header_row: int) -> dict[int, str]:
    return {
        col_idx: value_to_text(ws.cell(header_row, col_idx).value)
        for col_idx in range(1, ws.max_column + 1)
    }


def find_column(
    headers: dict[int, str],
    include_terms: list[str],
    exclude_terms: list[str] | None = None,
) -> int | None:
    include = [normalize(term) for term in include_terms]
    exclude = [normalize(term) for term in (exclude_terms or [])]

    for col_idx, header in headers.items():
        normalized = normalize(header)
        if not normalized:
            continue
        if all(term in normalized for term in include) and not any(term in normalized for term in exclude):
            return col_idx
    return None


def require_column(headers: dict[int, str], include_terms: list[str]) -> int:
    col_idx = find_column(headers, include_terms)
    if col_idx is None:
        raise ValueError(f"Column not found: {' '.join(include_terms)}")
    return col_idx


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def normalize_contextualization_subcategory_ids(data: Any) -> None:
    if not isinstance(data, dict):
        return

    for change in data.get("change_contextualizations", []):
        if not isinstance(change, dict):
            continue
        for field_name in ["framework_mappings", "prototype_actions"]:
            items = change.get(field_name, [])
            if not isinstance(items, list):
                continue
            for item in items:
                if not isinstance(item, dict):
                    continue
                if "subcategory_id" in item:
                    item["subcategory_id"] = canonicalize_subcategory_id(item["subcategory_id"])
                if "subcategory_text" in item and item["subcategory_text"] is not None:
                    item["subcategory_text"] = normalize_subcategory_references(item["subcategory_text"])
        control_actions = change.get("control_actions", [])
        if not isinstance(control_actions, list):
            continue
        for action in control_actions:
            if not isinstance(action, dict):
                continue
            if action.get("subcategories_with_weight") is not None:
                action["subcategories_with_weight"] = normalize_subcategory_references(
                    action["subcategories_with_weight"]
                )


def normalize_legacy_source_reference_fields(data: Any) -> None:
    """Accept contextualization JSON generated before source references were generalized."""
    if not isinstance(data, dict):
        return
    for change in data.get("change_contextualizations", []):
        if not isinstance(change, dict):
            continue
        for action in change.get("prototype_actions", []):
            if not isinstance(action, dict):
                continue
            if "source_reference" not in action and "circ285_reference" in action:
                action["source_reference"] = action.get("circ285_reference")
            action.pop("circ285_reference", None)


def remove_confidence_fields(data: Any) -> None:
    """Strip legacy and generated confidence fields from supplied JSON."""
    if not isinstance(data, dict):
        return
    data.pop("confidence_scoring", None)
    for change in data.get("change_contextualizations", []):
        if not isinstance(change, dict):
            continue
        for mapping in change.get("framework_mappings", []):
            if isinstance(mapping, dict):
                mapping.pop("confidence", None)
        for action_key in ["prototype_actions", "control_actions"]:
            for action in change.get(action_key, []):
                if isinstance(action, dict):
                    action.pop("confidence_assessment", None)


def collect_framework_subcategory_ids(framework_lookup: dict[str, dict[str, Any]]) -> list[str]:
    subcategory_ids = sorted(
        subcategory_id
        for subcategory_id in framework_lookup
        if value_to_text(subcategory_id)
    )
    if not subcategory_ids:
        raise ValueError("The framework core does not contain any usable subcategory IDs.")
    return subcategory_ids


def collect_workbook_subcategory_ids(bundle: dict[str, Any]) -> list[str]:
    subcategory_ids: set[str] = set()

    for row in bundle.get("contextualization_prototype_rows", []):
        if not isinstance(row, dict):
            continue
        subcategory_id = canonicalize_subcategory_id(row.get("subcategory_id"))
        if subcategory_id:
            subcategory_ids.add(subcategory_id)

    for row in bundle.get("existing_control_rows", []):
        if not isinstance(row, dict):
            continue
        for subcategory_id in subcategory_id_aliases_in_text(row.get("subcategories_with_weight")):
            subcategory_ids.add(canonicalize_subcategory_id(subcategory_id))

    return sorted(subcategory_ids)


def collect_expected_change_ids(bundle: dict[str, Any]) -> list[str]:
    change_ids = [
        value_to_text(change.get("change_id"))
        for change in bundle.get("annotated_changes", [])
        if isinstance(change, dict)
    ]
    missing = [index + 1 for index, change_id in enumerate(change_ids) if not change_id]
    duplicate_ids = sorted(
        change_id
        for change_id, count in Counter(change_ids).items()
        if change_id and count > 1
    )

    if missing or duplicate_ids:
        details: list[str] = []
        if missing:
            details.append(f"missing change_id values at positions: {', '.join(map(str, missing))}")
        if duplicate_ids:
            details.append(f"duplicate change_id values: {', '.join(duplicate_ids)}")
        raise ValueError("Invalid annotated changes: " + "; ".join(details))

    return change_ids


def validate_contextualization(
    data: dict[str, Any],
    output_schema: dict[str, Any],
    allowed_subcategory_ids: list[str],
    expected_change_ids: list[str],
    allowed_control_subcategory_ids: list[str] | None = None,
) -> None:
    errors: list[str] = []

    try:
        from jsonschema import Draft202012Validator
    except ImportError as error:
        print(
            "Error: jsonschema is required to validate contextualization output. "
            "Install the dependencies from requirements.txt.",
            file=sys.stderr,
        )
        raise SystemExit(1) from error

    validator = Draft202012Validator(output_schema)
    for error in sorted(validator.iter_errors(data), key=lambda item: list(item.path)):
        location = ".".join(str(part) for part in error.path) or "<root>"
        errors.append(f"{location}: {error.message}")

    change_contexts = data.get("change_contextualizations", []) if isinstance(data, dict) else []
    if not isinstance(change_contexts, list):
        change_contexts = []

    returned_change_ids = [
        value_to_text(change.get("change_id"))
        for change in change_contexts
        if isinstance(change, dict)
    ]
    returned_counts = Counter(returned_change_ids)
    expected_set = set(expected_change_ids)
    returned_set = set(returned_change_ids)

    missing_change_ids = sorted(expected_set - returned_set)
    unknown_change_ids = sorted(returned_set - expected_set)
    duplicate_change_ids = sorted(
        change_id
        for change_id, count in returned_counts.items()
        if change_id and count > 1
    )
    if len(change_contexts) != len(expected_change_ids):
        errors.append(
            "change_contextualizations must contain exactly one item for every annotated change "
            f"(expected {len(expected_change_ids)}, received {len(change_contexts)})"
        )
    if missing_change_ids:
        errors.append(f"missing change_id values: {', '.join(missing_change_ids)}")
    if unknown_change_ids:
        errors.append(f"unknown change_id values: {', '.join(unknown_change_ids)}")
    if duplicate_change_ids:
        errors.append(f"duplicate change_id values: {', '.join(duplicate_change_ids)}")

    allowed_subcategory_set = set(allowed_subcategory_ids)
    allowed_control_subcategory_set = allowed_subcategory_set | set(allowed_control_subcategory_ids or [])
    for change_index, change in enumerate(change_contexts):
        if not isinstance(change, dict):
            continue

        relevance = value_to_text(change.get("cybersecurity_relevance"))
        framework_mappings = change.get("framework_mappings", [])
        prototype_actions = change.get("prototype_actions", [])
        control_actions = change.get("control_actions", [])

        if relevance == "not_relevant":
            if value_to_text(change.get("classe")) != "Non applicabile":
                errors.append(
                    f"change_contextualizations.{change_index}.classe must be Non applicabile "
                    "when cybersecurity_relevance is not_relevant"
                )
            if value_to_text(change.get("priority")) != "Non applicabile":
                errors.append(
                    f"change_contextualizations.{change_index}.priority must be Non applicabile "
                    "when cybersecurity_relevance is not_relevant"
                )
            for field_name, value in [
                ("framework_mappings", framework_mappings),
                ("prototype_actions", prototype_actions),
                ("control_actions", control_actions),
            ]:
                if isinstance(value, list) and value:
                    errors.append(
                        f"change_contextualizations.{change_index}.{field_name} must be empty "
                        "when cybersecurity_relevance is not_relevant"
                    )

        for field_name, items in [
            ("framework_mappings", framework_mappings),
            ("prototype_actions", prototype_actions),
        ]:
            if not isinstance(items, list):
                continue
            for item_index, item in enumerate(items):
                if not isinstance(item, dict):
                    continue
                subcategory_id = value_to_text(item.get("subcategory_id"))
                if subcategory_id not in allowed_subcategory_set:
                    errors.append(
                        f"change_contextualizations.{change_index}.{field_name}.{item_index}.subcategory_id "
                        f"is not in the framework core: {subcategory_id or '<empty>'}"
                    )

        if not isinstance(control_actions, list):
            continue
        for action_index, action in enumerate(control_actions):
            if not isinstance(action, dict):
                continue
            if value_to_text(action.get("action")) == "new":
                if action.get("existing_control_row") is not None:
                    errors.append(
                        f"change_contextualizations.{change_index}.control_actions.{action_index}."
                        "existing_control_row must be null when action is new"
                    )
                if action.get("existing_control_id") is not None:
                    errors.append(
                        f"change_contextualizations.{change_index}.control_actions.{action_index}."
                        "existing_control_id must be null when action is new"
                    )
            for subcategory_id in SUBCATEGORY_RE.findall(
                value_to_text(action.get("subcategories_with_weight"))
            ):
                subcategory_id = canonicalize_subcategory_id(subcategory_id)
                if subcategory_id not in allowed_control_subcategory_set:
                    errors.append(
                        f"change_contextualizations.{change_index}.control_actions.{action_index}."
                        f"subcategories_with_weight contains an unknown framework/workbook subcategory: "
                        f"{subcategory_id}"
                    )

    if not errors:
        return

    print("Error: the contextualization JSON did not pass validation.", file=sys.stderr)
    for error in errors[:10]:
        print(f"- {error}", file=sys.stderr)
    if len(errors) > 10:
        print(f"... and {len(errors) - 10} more errors", file=sys.stderr)
    raise SystemExit(1)


def extract_framework_core(path: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    entries: list[dict[str, Any]] = []
    current_function = ""
    current_category = ""

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        function, category, subcategory, references = (list(row) + [None, None, None, None])[:4]
        if normalize(subcategory).startswith("subcategory"):
            continue
        if value_to_text(function):
            current_function = compact_text(function, 900)
        if value_to_text(category):
            current_category = compact_text(category, 900)
        if not value_to_text(subcategory):
            continue

        subcategory_text = compact_text(subcategory, 1200)
        subcategory_id = parse_subcategory_id(subcategory_text)
        if not SUBCATEGORY_ID_RE.fullmatch(subcategory_id):
            continue
        entry = {
            "row_number": row_idx,
            "function": current_function,
            "category": current_category,
            "subcategory_id": subcategory_id,
            "subcategory_text": subcategory_text,
            "informative_references": compact_text(references, 1200),
        }
        entries.append(entry)

    wb.close()
    return entries, {entry["subcategory_id"]: entry for entry in entries}


def extract_prototype(
    path: Path,
    prototype_reference_column: str | None,
) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb[wb.sheetnames[0]]
    header_row = find_header_row(ws, ["Subcategory", "Classe"])
    headers = header_map(ws, header_row)

    function_col = require_column(headers, ["function"])
    category_col = require_column(headers, ["category"])
    subcategory_col = require_column(headers, ["subcategory"])
    class_col = require_column(headers, ["classe"])
    priority_col = find_column(headers, ["priorit"]) or find_column(headers, ["livello"])
    source_reference_col = (
        parse_column_ref(prototype_reference_column)
        or find_column(headers, ["285"])
        or find_column(headers, ["circolare"])
    )

    rows: list[dict[str, Any]] = []
    current_function = ""
    current_category = ""

    for row_idx in range(header_row + 1, ws.max_row + 1):
        function = value_to_text(ws.cell(row_idx, function_col).value)
        category = value_to_text(ws.cell(row_idx, category_col).value)
        subcategory = value_to_text(ws.cell(row_idx, subcategory_col).value)

        if function:
            current_function = compact_text(function, 900)
        if category:
            current_category = compact_text(category, 900)
        if not subcategory:
            continue

        rows.append(
            {
                "row_number": row_idx,
                "function": current_function,
                "category": current_category,
                "subcategory_id": parse_subcategory_id(subcategory),
                "subcategory_text": compact_text(normalize_subcategory_references(subcategory), 1200),
                "classe": value_to_text(ws.cell(row_idx, class_col).value),
                "priority": value_to_text(ws.cell(row_idx, priority_col).value) if priority_col else "",
                "source_reference": (
                    compact_text(ws.cell(row_idx, source_reference_col).value, 1200)
                    if source_reference_col
                    else ""
                ),
            }
        )

    wb.close()
    return rows


def extract_controls(
    path: Path,
    controls_filter_column: str | None,
    controls_prototype_flag_column: str | None,
    controls_reference_pattern: str | None,
    circ285_only: bool,
    max_rows: int,
) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb[wb.sheetnames[0]]
    header_row = find_header_row(ws, ["Controlli", "Subcategories"])
    headers = header_map(ws, header_row)

    control_col = require_column(headers, ["controll"])
    subcategories_col = require_column(headers, ["subcategories"])
    metadata_headers = {column: header for column, header in headers.items()
                        if not normalize(header).startswith("ai ")}
    priority_col = find_column(metadata_headers, ["priorit"])
    delegated_to_col = find_column(metadata_headers, ["delegat"])
    general_reference_col = find_column(headers, ["riferimenti", "generale"])
    general_flag_col = find_column(headers, ["prototipo", "generale"], ["riferimenti"])
    filter_col = parse_column_ref(controls_filter_column) or general_reference_col
    target_flag_col = parse_column_ref(controls_prototype_flag_column) or general_flag_col
    reference_pattern = (
        re.compile(controls_reference_pattern, re.IGNORECASE)
        if controls_reference_pattern
        else (CIRC285_RE if circ285_only else None)
    )

    rows: list[dict[str, Any]] = []

    for row_idx in range(header_row + 1, ws.max_row + 1):
        control_text = value_to_text(ws.cell(row_idx, control_col).value)
        if not control_text:
            continue

        reference_value = value_to_text(ws.cell(row_idx, filter_col).value) if filter_col else ""
        general_reference = value_to_text(ws.cell(row_idx, general_reference_col).value) if general_reference_col else ""
        target_flag = value_to_text(ws.cell(row_idx, target_flag_col).value) if target_flag_col else ""

        is_selected_prototype = bool(reference_value or target_flag)
        if not is_selected_prototype:
            continue
        if reference_pattern and not reference_pattern.search(reference_value):
            continue

        rows.append(
            {
                "row_number": row_idx,
                "id": value_to_text(ws.cell(row_idx, 1).value),
                "control_text": compact_text(control_text, 0),
                "priority": (value_to_text(ws.cell(row_idx, priority_col).value) or None) if priority_col else None,
                "delegated_to": (value_to_text(ws.cell(row_idx, delegated_to_col).value) or None) if delegated_to_col else None,
                "subcategories_with_weight": compact_text(
                    normalize_subcategory_references(ws.cell(row_idx, subcategories_col).value),
                    1200,
                ),
                "general_references": compact_text(general_reference, 0),
                "target_references": compact_text(reference_value, 0),
                "filter_column": get_column_letter(filter_col) if filter_col else None,
                "filter_value": compact_text(reference_value, 1200),
                "prototype_flag_column": (
                    get_column_letter(target_flag_col) if target_flag_col else None
                ),
                "prototype_flag_value": compact_text(target_flag, 200),
            }
        )
        if max_rows and len(rows) >= max_rows:
            break

    wb.close()
    return rows


def compact_annotations(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    data = read_json(path)
    compact_changes: list[dict[str, Any]] = []

    for change in data.get("changes", []):
        raw_text = change.get("raw_text") or {}
        location = change.get("location") or {}
        compact_changes.append(
            {
                "change_id": change.get("change_id"),
                "summary": change.get("summary"),
                "target_location": location.get("target_location"),
                "classification": change.get("classification"),
                "classification_rationale": change.get("classification_rationale"),
                "target_text": compact_text(raw_text.get("target"), 0),
                "source_text": compact_text(raw_text.get("source"), 0),
                "references": change.get("references") or [],
                "notes": change.get("notes"),
            }
        )

    return data, compact_changes


def build_input_bundle(args: argparse.Namespace) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    annotations, compact_changes = compact_annotations(args.annotations_json)
    framework_core, framework_lookup = extract_framework_core(args.framework_core)
    legacy_circ285_filter = (
        not args.no_circ285_filter and not args.controls_reference_pattern
    )

    prototype_rows = extract_prototype(
        args.prototype_xlsx,
        args.prototype_reference_column,
    )

    control_rows = extract_controls(
        args.controls_xlsx,
        args.controls_filter_column,
        args.controls_prototype_flag_column,
        args.controls_reference_pattern,
        circ285_only=legacy_circ285_filter,
        max_rows=args.max_controls_context,
    )

    annotation_id = annotations.get("annotation_id") or args.annotations_json.stem
    document_id = annotations.get("document_id") or ""

    return (
        {
            "task": {
                "name": "update_contextualization",
                "language": args.output_language,
                "document_id": document_id,
                "source_annotation_id": annotation_id,
                "merge_policy": {
                    "default_mode": "review_copy_with_highlighting",
                    "do_not_delete_existing_rows": True,
                    "allow_actions": [
                        "new",
                        "modify",
                        "deprecate",
                        "no_change",
                        "review_needed",
                    ],
                    "non_cybersecurity_changes": "mark_not_relevant_and_leave_actions_empty",
                },
                "controls_filter": {
                    "circ285_only": legacy_circ285_filter,
                    "reference_column": args.controls_filter_column or "auto",
                    "prototype_flag_column": args.controls_prototype_flag_column or "auto",
                    "reference_pattern": args.controls_reference_pattern,
                },
            },
            "annotated_changes": compact_changes,
            "framework_core": framework_core,
            "contextualization_prototype_rows": prototype_rows,
            "existing_control_rows": control_rows,
        },
        framework_lookup,
    )


def run_query(args: argparse.Namespace, bundle_path: Path, schema_path: Path, output_json: Path) -> None:
    command = [
        sys.executable,
        str(args.query_script),
        "--prompt",
        str(args.prompt),
        "--schema",
        str(schema_path),
        "--file",
        str(bundle_path),
        "--output",
        str(output_json),
        "--final-reminder",
        contextualization_final_reminder(args.output_language),
    ]

    if args.llm:
        command[2:2] = ["--llm", args.llm]

    result = subprocess.run(command, text=True)
    if result.returncode != 0:
        raise SystemExit(result.returncode)


def run_self_consistency_check(
    args: argparse.Namespace,
    bundle_path: Path,
    schema_path: Path,
    output_schema: dict[str, Any],
    allowed_subcategory_ids: list[str],
    expected_change_ids: list[str],
    allowed_control_subcategory_ids: list[str],
) -> SelfConsistencyEvidence:
    if not args.confidence_self_consistency:
        return SelfConsistencyEvidence(False, {})

    independent_output = args.output_dir / "contextualization_self_consistency.json"
    print("Running optional independent contextualization pass for confidence scoring...")
    try:
        run_query(args, bundle_path, schema_path, independent_output)
        independent_data = read_json(independent_output)
        normalize_contextualization_subcategory_ids(independent_data)
        normalize_legacy_source_reference_fields(independent_data)
        validate_contextualization(
            independent_data,
            output_schema,
            allowed_subcategory_ids,
            expected_change_ids,
            allowed_control_subcategory_ids,
        )
    except SystemExit:
        print(
            "Warning: the optional independent pass was unavailable; no self-consistency "
            "bonus will be applied.",
            file=sys.stderr,
        )
        return SelfConsistencyEvidence(True, {})

    print(f"Independent contextualization JSON: {independent_output}")
    return build_self_consistency_evidence(independent_data)


def action_fill(action: str) -> PatternFill:
    return PatternFill("solid", fgColor=ACTION_FILL_COLORS.get(action, "DDEBF7"))


def display_action(action: Any) -> str:
    action_text = value_to_text(action)
    return ACTION_LABELS.get(action_text, action_text)


def display_merge_action(original_action: str, effective_action: str) -> str:
    if original_action == effective_action:
        return display_action(original_action)
    return f"{display_action(effective_action)} ({display_action(original_action)})"


def effective_merge_action(action_name: str, requires_review: bool) -> str:
    return "review_needed" if requires_review else action_name


def confidence_aware_merge_action(action: dict[str, Any], effective_action: str) -> str:
    assessment = action.get("confidence_assessment")
    if isinstance(assessment, dict) and assessment.get("level") == "review_needed":
        return "review_needed"
    return effective_action


def add_merge_review_event(
    events: list[MergeReviewEvent],
    change_id: str,
    workbook: str,
    original_action: str,
    effective_action: str,
    row_idx: int | None,
    subject: Any,
    review_status: str,
) -> None:
    events.append(
        MergeReviewEvent(
            change_id=change_id,
            workbook=workbook,
            original_action=original_action,
            effective_action=effective_action,
            row_idx=row_idx,
            subject=value_to_text(subject),
            review_status=review_status,
        )
    )


def display_relevance(relevance: Any) -> str:
    relevance_text = value_to_text(relevance)
    return RELEVANCE_LABELS.get(relevance_text, relevance_text)


def new_merge_stats() -> dict[str, set[int]]:
    return {
        "new_rows": set(),
        "modified_rows": set(),
        "deprecated_rows": set(),
        "review_needed_rows": set(),
    }


def record_merged_row(stats: dict[str, set[int]], action: str, row_idx: int) -> None:
    statistic_key = MERGE_STAT_KEYS.get(action)
    if statistic_key:
        if action == "review_needed":
            for key in MERGE_STAT_KEYS.values():
                if key != statistic_key:
                    stats[key].discard(row_idx)
        stats[statistic_key].add(row_idx)


def effective_row_action(
    row_actions: dict[int, str],
    row_idx: int,
    action_name: str,
    requires_review: bool,
) -> str:
    action = effective_merge_action(action_name, requires_review)
    if row_actions.get(row_idx) == "review_needed":
        action = "review_needed"
    row_actions[row_idx] = action
    return action


def serialize_merge_stats(stats: dict[str, set[int]]) -> dict[str, int]:
    return {key: len(rows) for key, rows in stats.items()}


def copy_cell_format(source: Any, target: Any) -> None:
    if not source.has_style:
        return
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy.copy(source.protection)


def copy_row_format(ws: Any, source_row: int | None, target_row: int) -> None:
    if not source_row or source_row < 1 or source_row > ws.max_row:
        return

    source_dimension = ws.row_dimensions[source_row]
    target_dimension = ws.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed

    for col_idx in range(1, ws.max_column + 1):
        copy_cell_format(ws.cell(source_row, col_idx), ws.cell(target_row, col_idx))


def find_last_data_row(ws: Any, key_col: int, header_row: int) -> int | None:
    for row_idx in range(ws.max_row, header_row, -1):
        if value_to_text(ws.cell(row_idx, key_col).value):
            return row_idx
    return None


def find_last_value_row(ws: Any, header_row: int, max_col: int | None = None) -> int:
    last_col = max_col or ws.max_column
    for row_idx in range(ws.max_row, header_row, -1):
        for col_idx in range(1, last_col + 1):
            if value_to_text(ws.cell(row_idx, col_idx).value):
                return row_idx
    return header_row


def set_review_alignment(cell: Any) -> None:
    alignment = copy.copy(cell.alignment)
    alignment.wrap_text = True
    alignment.vertical = "top"
    cell.alignment = alignment


def style_review_header(cell: Any) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    set_review_alignment(cell)


def align_review_cells(ws: Any) -> None:
    for row in ws.iter_rows():
        for cell in row:
            set_review_alignment(cell)


def apply_row_fill(
    ws: Any,
    row_idx: int,
    fill: PatternFill,
    excluded_columns: Iterable[int] = (),
) -> None:
    excluded = set(excluded_columns)
    for col_idx in range(1, ws.max_column + 1):
        if col_idx in excluded:
            continue
        ws.cell(row_idx, col_idx).fill = copy.copy(fill)


def clear_column_borders(ws: Any, col_indices: Iterable[int], header_row: int) -> None:
    for col_idx in col_indices:
        for row_idx in range(header_row, ws.max_row + 1):
            ws.cell(row_idx, col_idx).border = Border()


def format_separator_column(ws: Any, col_idx: int, header_row: int) -> None:
    ws.column_dimensions[get_column_letter(col_idx)].width = 6
    for row_idx in range(header_row, ws.max_row + 1):
        cell = ws.cell(row_idx, col_idx)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()


def is_blank_column(ws: Any, col_idx: int, header_row: int) -> bool:
    if col_idx < 1:
        return False
    return all(
        not value_to_text(ws.cell(row_idx, col_idx).value)
        for row_idx in range(header_row, ws.max_row + 1)
    )


def extend_column_format_from_left(ws: Any, col_idx: int, header_row: int) -> None:
    if col_idx <= 1:
        return
    for row_idx in range(header_row + 1, ws.max_row + 1):
        copy_cell_format(ws.cell(row_idx, col_idx - 1), ws.cell(row_idx, col_idx))


def ensure_columns(
    ws: Any,
    headers: list[str],
    header_row: int = 1,
    separator_before: bool = False,
    extend_body_format: bool = True,
    borderless: bool = False,
) -> dict[str, int]:
    existing = {
        value_to_text(ws.cell(header_row, col_idx).value): col_idx
        for col_idx in range(1, ws.max_column + 1)
    }
    result: dict[str, int] = {}
    missing_headers = [header for header in headers if header not in existing]

    if separator_before and missing_headers:
        format_separator_column(ws, ws.max_column + 1, header_row)

    for header in headers:
        if header in existing:
            result[header] = existing[header]
            continue

        col_idx = ws.max_column + 1
        cell = ws.cell(header_row, col_idx)
        cell.value = header
        style_review_header(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(header) + 2, 14), 38)
        if extend_body_format:
            extend_column_format_from_left(ws, col_idx, header_row)
        if borderless:
            clear_column_borders(ws, {col_idx}, header_row)
        result[header] = col_idx

    return result


def append_ai_value(cell: Any, value: Any) -> None:
    text = compact_text(value, 3000) or "-"
    existing = value_to_text(cell.value)
    cell.value = f"{existing}\n{text}" if existing else text


def append_ai_values(
    ws: Any,
    row_idx: int,
    ai_columns: dict[str, int],
    values: dict[str, Any],
) -> None:
    for header, value in values.items():
        append_ai_value(ws.cell(row_idx, ai_columns[header]), value)


def prototype_ai_values(
    action: dict[str, Any],
    change_id: str,
    review_status: str,
    effective_action: str,
) -> dict[str, Any]:
    return {
        "AI Azione": display_merge_action(value_to_text(action.get("action")), effective_action),
        "AI Change ID": change_id,
        "AI Punteggio confidenza": confidence_display(action.get("confidence_assessment")),
        "AI Classe proposta": action.get("classe"),
        "AI Priorità proposta": action.get("priority"),
        "AI Riferimento fonte": action.get("source_reference"),
        "AI Razionale": action.get("rationale"),
        "AI Stato revisione": review_status,
    }


def control_ai_values(
    action: dict[str, Any],
    change_id: str,
    review_status: str,
    effective_action: str,
) -> dict[str, Any]:
    return {
        "AI Azione": display_merge_action(value_to_text(action.get("action")), effective_action),
        "AI Change ID": change_id,
        "AI Punteggio confidenza": confidence_display(action.get("confidence_assessment")),
        "AI Controllo proposto": action.get("control_text"),
        "AI Subcategories proposte": action.get("subcategories_with_weight"),
        "AI Riferimenti proposti": action.get("references"),
        "AI Priorità": action.get("priority"),
        "AI Delegato a": action.get("delegated_to"),
        "AI Razionale": action.get("rationale"),
        "AI Stato revisione": review_status,
    }


def writable_comment_cell(cell: Any) -> Any | None:
    if cell.__class__.__name__ != "MergedCell":
        return cell

    ws = cell.parent
    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return ws.cell(merged_range.min_row, merged_range.min_col)
    return None


def add_comment(cell: Any, text: str) -> None:
    text = value_to_text(text)
    if not text:
        return
    cell = writable_comment_cell(cell)
    if cell is None:
        return
    existing = cell.comment.text if cell.comment else ""
    combined = f"{existing}\n\n{text}".strip() if existing else text
    cell.comment = Comment(combined[:3000], "AI")


def find_prototype_columns(
    ws: Any,
    prototype_reference_column: str | None = None,
) -> dict[str, int]:
    header_row = find_header_row(ws, ["Subcategory", "Classe"])
    headers = header_map(ws, header_row)
    return {
        "header_row": header_row,
        "function": require_column(headers, ["function"]),
        "category": require_column(headers, ["category"]),
        "subcategory": require_column(headers, ["subcategory"]),
        "classe": require_column(headers, ["classe"]),
        "priority": find_column(headers, ["priorit"]) or find_column(headers, ["livello"]),
        "source_reference": (
            parse_column_ref(prototype_reference_column)
            or find_column(headers, ["285"])
            or find_column(headers, ["circolare"])
        ),
    }


def find_controls_columns(
    ws: Any,
    controls_filter_column: str | None,
    controls_prototype_flag_column: str | None = None,
) -> dict[str, int | None]:
    header_row = find_header_row(ws, ["Controlli", "Subcategories"])
    headers = header_map(ws, header_row)
    general_reference_col = find_column(headers, ["riferimenti", "generale"])
    return {
        "header_row": header_row,
        "id": find_column(headers, ["id"]) or 1,
        "control": require_column(headers, ["controll"]),
        "subcategories": require_column(headers, ["subcategories"]),
        "general_reference": general_reference_col,
        "target_flag": (
            parse_column_ref(controls_prototype_flag_column)
            or find_column(headers, ["prototipo", "generale"], ["riferimenti"])
        ),
        "target_reference": parse_column_ref(controls_filter_column) or general_reference_col,
        "priority": find_column(headers, ["priorit"]),
        "delegated_to": find_column(headers, ["delegat"]),
    }


def index_prototype_rows(ws: Any, subcategory_col: int, header_row: int) -> dict[str, int]:
    index: dict[str, int] = {}
    for row_idx in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row_idx, subcategory_col).value
        if value:
            for subcategory_id in subcategory_id_aliases(value):
                if subcategory_id not in index:
                    index[subcategory_id] = row_idx
    return index


def index_control_rows(
    ws: Any,
    subcategories_col: int,
    header_row: int,
) -> dict[str, set[int]]:
    index: dict[str, set[int]] = {}
    for row_idx in range(header_row + 1, ws.max_row + 1):
        for subcategory_id in subcategory_id_aliases_in_text(
            ws.cell(row_idx, subcategories_col).value
        ):
            index.setdefault(subcategory_id, set()).add(row_idx)
    return index


def resolve_control_row(
    index: dict[str, set[int]],
    subcategories_with_weight: Any,
) -> int | None:
    candidates: set[int] = set()
    for subcategory_id in subcategory_id_aliases_in_text(subcategories_with_weight):
        candidates.update(index.get(subcategory_id, set()))
    if len(candidates) == 1:
        return next(iter(candidates))
    return None


def integer_row(value: Any) -> int | None:
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def subcategory_ids_from_text(value: Any) -> set[str]:
    subcategory_ids = {
        canonicalize_subcategory_id(subcategory_id)
        for subcategory_id in subcategory_id_aliases_in_text(value)
    }
    if subcategory_ids:
        return {subcategory_id for subcategory_id in subcategory_ids if subcategory_id}

    parsed = parse_subcategory_id(value)
    if SUBCATEGORY_ID_RE.fullmatch(parsed):
        return {canonicalize_subcategory_id(parsed)}
    return set()


def format_subcategory_ids(subcategory_ids: set[str]) -> str:
    if not subcategory_ids:
        return "<empty>"
    return ", ".join(sorted(subcategory_ids))


def resolve_prototype_target_row(
    ws: Any,
    subcategory_col: int,
    row_index: dict[str, int],
    supplied_row_value: Any,
    subcategory_id: str,
) -> RowResolution:
    fallback_row = row_index.get(subcategory_id)
    supplied_row = integer_row(supplied_row_value)

    if supplied_row is not None and 1 <= supplied_row <= ws.max_row:
        actual_ids = subcategory_ids_from_text(ws.cell(supplied_row, subcategory_col).value)
        if subcategory_id in actual_ids:
            return RowResolution(supplied_row, "Da revisionare")

        mismatch = (
            f"existing_row AI {supplied_row} ignorata: "
            f"riga con {format_subcategory_ids(actual_ids)}, atteso {subcategory_id}"
        )
        if fallback_row:
            return RowResolution(
                fallback_row,
                f"Da revisionare ({mismatch}; usata riga {fallback_row})",
                True,
            )
        return RowResolution(None, f"Proposta senza riga esistente coerente ({mismatch})", True)

    if supplied_row_value is not None:
        invalid = (
            f"existing_row AI non valida: {supplied_row_value}"
            if supplied_row is None
            else f"existing_row AI fuori intervallo: {supplied_row}"
        )
        if fallback_row:
            return RowResolution(
                fallback_row,
                f"Da revisionare ({invalid}; usata riga {fallback_row})",
                True,
            )
        return RowResolution(None, f"Proposta senza riga esistente coerente ({invalid})", True)

    if fallback_row:
        return RowResolution(fallback_row, "Da revisionare")
    return RowResolution(None, "Proposta senza riga esistente corrispondente", True)


def resolve_controls_target_row(
    ws: Any,
    subcategories_col: int,
    row_index: dict[str, set[int]],
    supplied_row_value: Any,
    subcategories_with_weight: Any,
) -> RowResolution:
    proposed_ids = subcategory_ids_from_text(subcategories_with_weight)
    fallback_row = resolve_control_row(row_index, subcategories_with_weight)
    supplied_row = integer_row(supplied_row_value)

    if supplied_row is not None and 1 <= supplied_row <= ws.max_row:
        actual_ids = subcategory_ids_from_text(ws.cell(supplied_row, subcategories_col).value)
        if proposed_ids and actual_ids.intersection(proposed_ids):
            return RowResolution(supplied_row, "Da revisionare")

        mismatch = (
            f"existing_control_row AI {supplied_row} ignorata: "
            f"riga con {format_subcategory_ids(actual_ids)}, "
            f"proposta {format_subcategory_ids(proposed_ids)}"
        )
        if fallback_row:
            return RowResolution(
                fallback_row,
                f"Da revisionare ({mismatch}; usata riga {fallback_row})",
                True,
            )
        return RowResolution(None, f"Proposta senza riga controllo coerente ({mismatch})", True)

    if supplied_row_value is not None:
        invalid = (
            f"existing_control_row AI non valida: {supplied_row_value}"
            if supplied_row is None
            else f"existing_control_row AI fuori intervallo: {supplied_row}"
        )
        if fallback_row:
            return RowResolution(
                fallback_row,
                f"Da revisionare ({invalid}; usata riga {fallback_row})",
                True,
            )
        return RowResolution(None, f"Proposta senza riga controllo coerente ({invalid})", True)

    if fallback_row:
        return RowResolution(fallback_row, "Da revisionare")
    return RowResolution(None, "Proposta senza riga controllo corrispondente", True)


def iter_change_contexts(data: dict[str, Any]) -> list[dict[str, Any]]:
    return data.get("change_contextualizations", [])


def annotate_confidence_assessments(
    data: dict[str, Any],
    prototype_path: Path,
    controls_path: Path,
    framework_lookup: dict[str, dict[str, Any]],
    controls_filter_column: str | None,
    controls_prototype_flag_column: str | None,
    prototype_reference_column: str | None,
    self_consistency: SelfConsistencyEvidence,
) -> None:
    """Attach deterministic confidence to each action that may change a workbook."""
    prototype_wb = load_workbook(prototype_path, read_only=True, data_only=True)
    prototype_ws = prototype_wb[prototype_wb.sheetnames[0]]
    prototype_columns = find_prototype_columns(
        prototype_ws,
        prototype_reference_column,
    )
    prototype_row_index = index_prototype_rows(
        prototype_ws,
        prototype_columns["subcategory"],
        prototype_columns["header_row"],
    )

    controls_wb = load_workbook(controls_path, read_only=True, data_only=True)
    controls_ws = controls_wb[controls_wb.sheetnames[0]]
    controls_columns = find_controls_columns(
        controls_ws,
        controls_filter_column,
        controls_prototype_flag_column,
    )
    controls_row_index = index_control_rows(
        controls_ws,
        int(controls_columns["subcategories"]),
        int(controls_columns["header_row"]),
    )

    try:
        for change in iter_change_contexts(data):
            for action in change.get("prototype_actions", []):
                if not isinstance(action, dict) or value_to_text(action.get("action")) == "no_change":
                    continue
                action_name = value_to_text(action.get("action"))
                subcategory_id = canonicalize_subcategory_id(action.get("subcategory_id"))
                if action_name == "new":
                    target_validation_passed = (
                        action.get("existing_row") is None
                        and bool(subcategory_id)
                        and subcategory_id not in prototype_row_index
                    )
                    resolution = RowResolution(
                        None,
                        "Nuova riga proposta"
                        if target_validation_passed
                        else "Nuova riga proposta per una sottocategoria già presente",
                        not target_validation_passed,
                    )
                else:
                    resolution = resolve_prototype_target_row(
                        prototype_ws,
                        prototype_columns["subcategory"],
                        prototype_row_index,
                        action.get("existing_row"),
                        subcategory_id,
                    )
                    target_validation_passed = (
                        resolution.row_idx is not None and not resolution.requires_review
                    )
                action["confidence_assessment"] = score_action_confidence(
                    change,
                    action,
                    "prototype",
                    framework_lookup,
                    target_validation_passed,
                    resolution.requires_review,
                    effective_merge_action(action_name, resolution.requires_review),
                    self_consistency,
                )

            for action in change.get("control_actions", []):
                if not isinstance(action, dict) or value_to_text(action.get("action")) == "no_change":
                    continue
                action_name = value_to_text(action.get("action"))
                if action_name == "new":
                    proposed_subcategory_ids = subcategory_ids_from_text(
                        action.get("subcategories_with_weight")
                    )
                    target_validation_passed = (
                        action.get("existing_control_row") is None
                        and bool(proposed_subcategory_ids)
                        and not any(
                            controls_row_index.get(subcategory_id)
                            for subcategory_id in proposed_subcategory_ids
                        )
                    )
                    resolution = RowResolution(
                        None,
                        "Nuovo controllo proposto"
                        if target_validation_passed
                        else "Nuovo controllo proposto con sottocategorie già coperte",
                        not target_validation_passed,
                    )
                else:
                    resolution = resolve_controls_target_row(
                        controls_ws,
                        int(controls_columns["subcategories"]),
                        controls_row_index,
                        action.get("existing_control_row"),
                        action.get("subcategories_with_weight"),
                    )
                    target_validation_passed = (
                        resolution.row_idx is not None and not resolution.requires_review
                    )
                action["confidence_assessment"] = score_action_confidence(
                    change,
                    action,
                    "control",
                    framework_lookup,
                    target_validation_passed,
                    resolution.requires_review,
                    effective_merge_action(action_name, resolution.requires_review),
                    self_consistency,
                )
    finally:
        prototype_wb.close()
        controls_wb.close()

    data["confidence_scoring"] = {
        "method": "deterministic_validation_plus_optional_independent_run_agreement",
        "version": "1.0",
        "self_consistency_requested": self_consistency.requested,
        "weights": CONFIDENCE_WEIGHTS,
    }


def write_review_sheet(
    wb: Any,
    data: dict[str, Any],
    merge_stats: dict[str, int],
    review_events: list[MergeReviewEvent],
) -> None:
    for sheet_name in ["AI Change Log", "Registro modifiche AI"]:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
    sheet_name = "Registro modifiche AI"
    ws = wb.create_sheet(sheet_name)

    ws.append(["Riepilogo aggiornamento", "Numero di righe"])
    for cell in ws[1]:
        style_review_header(cell)
    for label, key in MERGE_STAT_LABELS:
        ws.append([label, merge_stats[key]])
    ws.append([])

    headers = [
        "ID modifica",
        "Rilevanza",
        "Classe",
        "Priorità",
        "Sottocategoria",
        "Azioni prototipo",
        "Azioni controlli",
        "Confidenza",
        "Motivazione",
        "Note",
    ]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        style_review_header(cell)

    for change in iter_change_contexts(data):
        mappings = ", ".join(
            f"{item.get('subcategory_id')} ({item.get('mapping_weight')})"
            for item in change.get("framework_mappings", [])
        )
        prototype_actions = ", ".join(
            f"{display_action(item.get('action'))}:{item.get('subcategory_id')}"
            for item in change.get("prototype_actions", [])
        )
        control_actions = ", ".join(
            f"{display_action(item.get('action'))}:{item.get('existing_control_row') or 'nuovo'}"
            for item in change.get("control_actions", [])
        )
        ws.append(
            [
                change.get("change_id"),
                display_relevance(change.get("cybersecurity_relevance")),
                change.get("classe"),
                change.get("priority"),
                mappings,
                prototype_actions,
                control_actions,
                change_confidence_summary(change),
                change.get("relevance_rationale"),
                change.get("notes"),
            ]
        )

    if review_events:
        ws.append([])
        ws.append(
            [
                "Esiti di merge da revisionare",
                "Ambito",
                "Azione proposta",
                "Esito applicato",
                "Riga",
                "Oggetto",
                "Stato revisione",
            ]
        )
        for cell in ws[ws.max_row]:
            style_review_header(cell)
        for event in review_events:
            ws.append(
                [
                    event.change_id,
                    event.workbook,
                    display_action(event.original_action),
                    display_merge_action(event.original_action, event.effective_action),
                    event.row_idx,
                    event.subject,
                    event.review_status,
                ]
            )

    for col_idx, width in enumerate([28, 18, 18, 14, 36, 36, 36, 34, 70, 40], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    align_review_cells(ws)


def finish_workbook_merge(
    wb: Any,
    ws: Any,
    data: dict[str, Any],
    stats: dict[str, set[int]],
    review_events: list[MergeReviewEvent],
    output_path: Path,
) -> dict[str, int]:
    align_review_cells(ws)
    merge_stats = serialize_merge_stats(stats)
    write_review_sheet(wb, data, merge_stats, review_events)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return merge_stats


def has_generated_action_fill(cell: Any) -> bool:
    fill = cell.fill
    color = value_to_text(getattr(fill.fgColor, "rgb", "")).upper()
    return (
        fill.fill_type == "solid"
        and len(color) >= 6
        and color[-6:] in ACTION_FILL_COLORS.values()
    )


def remove_clean_presentation_artifacts(wb: Any, ws: Any, header_row: int) -> None:
    """Remove reviewer-only AI columns, comments, highlights, and change-log sheets."""
    for sheet_name in ["AI Change Log", "Registro modifiche AI"]:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    ai_columns = [
        col_idx
        for col_idx in range(1, ws.max_column + 1)
        if value_to_text(ws.cell(header_row, col_idx).value).startswith("AI ")
    ]
    if ai_columns:
        first_ai_column = min(ai_columns)
        columns_to_remove = set(ai_columns)
        separator_column = first_ai_column - 1
        if is_blank_column(ws, separator_column, header_row):
            columns_to_remove.add(separator_column)
        for col_idx in sorted(columns_to_remove, reverse=True):
            ws.delete_cols(col_idx)

    for row in ws.iter_rows():
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            cell.comment = None
            if has_generated_action_fill(cell):
                cell.fill = PatternFill(fill_type=None)


def set_if_not_none(cell: Any, value: Any) -> None:
    if value is not None:
        cell.value = value


def merge_clean_prototype(
    prototype_path: Path,
    output_path: Path,
    data: dict[str, Any],
    framework_lookup: dict[str, dict[str, Any]],
    prototype_reference_column: str | None,
) -> dict[str, int]:
    """Create a publication-ready prototype workbook from unambiguous actions."""
    wb = load_workbook(prototype_path)
    ws = wb[wb.sheetnames[0]]
    columns = find_prototype_columns(ws, prototype_reference_column)
    remove_clean_presentation_artifacts(wb, ws, columns["header_row"])
    columns = find_prototype_columns(ws, prototype_reference_column)

    row_index = index_prototype_rows(ws, columns["subcategory"], columns["header_row"])
    template_row = find_last_data_row(ws, columns["subcategory"], columns["header_row"])
    next_row = find_last_value_row(ws, columns["header_row"]) + 1
    stats = new_merge_stats()
    review_needed_targets = review_needed_target_keys(data, "prototype")

    for change in iter_change_contexts(data):
        for action in change.get("prototype_actions", []):
            if not clean_action_is_publishable(action) or clean_action_conflicts_with_review(
                action, "prototype", review_needed_targets
            ):
                continue
            action_name = value_to_text(action.get("action"))
            subcategory_id = canonicalize_subcategory_id(action.get("subcategory_id"))

            if action_name == "new":
                row_idx = next_row
                next_row += 1
                framework_row = framework_lookup.get(subcategory_id, {})
                copy_row_format(ws, template_row, row_idx)
                ws.cell(row_idx, columns["function"]).value = framework_row.get("function")
                ws.cell(row_idx, columns["category"]).value = framework_row.get("category")
                ws.cell(row_idx, columns["subcategory"]).value = (
                    action.get("subcategory_text")
                    or framework_row.get("subcategory_text")
                    or subcategory_id
                )
                set_if_not_none(ws.cell(row_idx, columns["classe"]), action.get("classe"))
                if columns["priority"]:
                    set_if_not_none(ws.cell(row_idx, columns["priority"]), action.get("priority"))
                if columns["source_reference"]:
                    set_if_not_none(
                        ws.cell(row_idx, columns["source_reference"]),
                        action.get("source_reference"),
                    )
                record_merged_row(stats, action_name, row_idx)
                if subcategory_id:
                    row_index.setdefault(subcategory_id, row_idx)
                continue

            if action_name != "modify":
                continue

            resolution = resolve_prototype_target_row(
                ws,
                columns["subcategory"],
                row_index,
                action.get("existing_row"),
                subcategory_id,
            )
            if resolution.requires_review or resolution.row_idx is None:
                continue

            row_idx = resolution.row_idx
            set_if_not_none(
                ws.cell(row_idx, columns["subcategory"]),
                action.get("subcategory_text"),
            )
            set_if_not_none(ws.cell(row_idx, columns["classe"]), action.get("classe"))
            if columns["priority"]:
                set_if_not_none(ws.cell(row_idx, columns["priority"]), action.get("priority"))
            if columns["source_reference"]:
                set_if_not_none(
                    ws.cell(row_idx, columns["source_reference"]),
                    action.get("source_reference"),
                )
            record_merged_row(stats, action_name, row_idx)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return serialize_merge_stats(stats)


def merge_clean_controls(
    controls_path: Path,
    output_path: Path,
    data: dict[str, Any],
    controls_filter_column: str | None,
    controls_prototype_flag_column: str | None,
) -> dict[str, int]:
    """Create a publication-ready controls workbook from unambiguous actions."""
    wb = load_workbook(controls_path)
    ws = wb[wb.sheetnames[0]]
    columns = find_controls_columns(
        ws,
        controls_filter_column,
        controls_prototype_flag_column,
    )
    header_row = int(columns["header_row"])
    remove_clean_presentation_artifacts(wb, ws, header_row)
    columns = find_controls_columns(
        ws,
        controls_filter_column,
        controls_prototype_flag_column,
    )
    header_row = int(columns["header_row"])

    row_index = index_control_rows(ws, int(columns["subcategories"]), header_row)
    template_row = find_last_data_row(ws, int(columns["control"]), header_row)
    next_row = find_last_value_row(ws, header_row) + 1
    stats = new_merge_stats()
    review_needed_targets = review_needed_target_keys(data, "control")

    for change in iter_change_contexts(data):
        for action in change.get("control_actions", []):
            if not clean_action_is_publishable(action) or clean_action_conflicts_with_review(
                action, "control", review_needed_targets
            ):
                continue
            action_name = value_to_text(action.get("action"))

            if action_name == "new":
                row_idx = next_row
                next_row += 1
                copy_row_format(ws, template_row, row_idx)
                id_column = int(columns["id"] or 1)
                ws.cell(row_idx, id_column).value = f"=ROW(A{row_idx - 1})"
            elif action_name == "modify":
                resolution = resolve_controls_target_row(
                    ws,
                    int(columns["subcategories"]),
                    row_index,
                    action.get("existing_control_row"),
                    action.get("subcategories_with_weight"),
                )
                if resolution.requires_review or resolution.row_idx is None:
                    continue
                row_idx = resolution.row_idx
            else:
                continue

            set_if_not_none(ws.cell(row_idx, int(columns["control"])), action.get("control_text"))
            set_if_not_none(
                ws.cell(row_idx, int(columns["subcategories"])),
                action.get("subcategories_with_weight"),
            )
            target_reference = columns.get("target_reference")
            if target_reference:
                set_if_not_none(ws.cell(row_idx, int(target_reference)), action.get("references"))
            priority_column = columns.get("priority")
            if priority_column:
                set_if_not_none(ws.cell(row_idx, int(priority_column)), action.get("priority"))
            delegated_to_column = columns.get("delegated_to")
            if delegated_to_column:
                set_if_not_none(
                    ws.cell(row_idx, int(delegated_to_column)),
                    action.get("delegated_to"),
                )
            if action_name == "new":
                target_flag = columns.get("target_flag")
                if target_flag:
                    ws.cell(row_idx, int(target_flag)).value = "X "
            record_merged_row(stats, action_name, row_idx)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return serialize_merge_stats(stats)


def merge_prototype(
    prototype_path: Path,
    output_path: Path,
    data: dict[str, Any],
    framework_lookup: dict[str, dict[str, Any]],
    prototype_reference_column: str | None,
) -> dict[str, int]:
    wb = load_workbook(prototype_path)
    ws = wb[wb.sheetnames[0]]
    columns = find_prototype_columns(ws, prototype_reference_column)
    main_table_last_col = ws.max_column
    ai_cols = ensure_columns(
        ws,
        [
            "AI Azione",
            "AI Change ID",
            "AI Punteggio confidenza",
            "AI Classe proposta",
            "AI Priorità proposta",
            "AI Riferimento fonte",
            "AI Razionale",
            "AI Stato revisione",
        ],
        header_row=columns["header_row"],
    )
    row_index = index_prototype_rows(ws, columns["subcategory"], columns["header_row"])
    prototype_template_row = find_last_data_row(ws, columns["subcategory"], columns["header_row"])
    prototype_next_row = find_last_value_row(
        ws,
        columns["header_row"],
        main_table_last_col,
    ) + 1
    stats = new_merge_stats()
    review_events: list[MergeReviewEvent] = []
    row_actions: dict[int, str] = {}

    for change in iter_change_contexts(data):
        change_id = value_to_text(change.get("change_id"))
        for action in change.get("prototype_actions", []):
            action_name = value_to_text(action.get("action"))
            if action_name == "no_change":
                continue

            subcategory_id = canonicalize_subcategory_id(action.get("subcategory_id"))
            resolution = (
                RowResolution(None, "Nuova riga proposta")
                if action_name == "new"
                else resolve_prototype_target_row(
                    ws,
                    columns["subcategory"],
                    row_index,
                    action.get("existing_row"),
                    subcategory_id,
                )
            )
            effective_action = confidence_aware_merge_action(
                action,
                effective_merge_action(action_name, resolution.requires_review),
            )

            if resolution.row_idx and 1 <= resolution.row_idx <= ws.max_row:
                row_idx = resolution.row_idx
                effective_action = confidence_aware_merge_action(
                    action,
                    effective_row_action(
                        row_actions,
                        row_idx,
                        action_name,
                        resolution.requires_review,
                    ),
                )
                append_ai_values(
                    ws,
                    row_idx,
                    ai_cols,
                    prototype_ai_values(
                        action,
                        change_id,
                        resolution.review_status,
                        effective_action,
                    ),
                )
                apply_row_fill(
                    ws,
                    row_idx,
                    action_fill(effective_action),
                    excluded_columns=(columns["function"],),
                )
                add_comment(ws.cell(row_idx, columns["subcategory"]), action.get("rationale") or "")
                record_merged_row(stats, effective_action, row_idx)
                if resolution.requires_review:
                    add_merge_review_event(
                        review_events,
                        change_id,
                        "Prototipo",
                        action_name,
                        effective_action,
                        row_idx,
                        subcategory_id,
                        resolution.review_status,
                    )
                continue

            if action_name not in APPENDABLE_ACTIONS:
                if action_name == "deprecate":
                    add_merge_review_event(
                        review_events,
                        change_id,
                        "Prototipo",
                        action_name,
                        effective_action,
                        None,
                        subcategory_id,
                        f"Deprecazione non applicata: {resolution.review_status}",
                    )
                continue

            framework_row = framework_lookup.get(subcategory_id, {})
            new_row = prototype_next_row
            prototype_next_row += 1
            effective_action = confidence_aware_merge_action(
                action,
                effective_row_action(
                    row_actions,
                    new_row,
                    action_name,
                    resolution.requires_review,
                ),
            )
            copy_row_format(ws, prototype_template_row, new_row)
            ws.cell(new_row, columns["function"]).value = framework_row.get("function")
            ws.cell(new_row, columns["category"]).value = framework_row.get("category")
            ws.cell(new_row, columns["subcategory"]).value = (
                action.get("subcategory_text") or framework_row.get("subcategory_text") or subcategory_id
            )
            ws.cell(new_row, columns["classe"]).value = action.get("classe")
            if columns["priority"]:
                ws.cell(new_row, columns["priority"]).value = action.get("priority")
            if columns["source_reference"]:
                ws.cell(new_row, columns["source_reference"]).value = action.get(
                    "source_reference"
                )
            append_ai_values(
                ws,
                new_row,
                ai_cols,
                prototype_ai_values(
                    action,
                    change_id,
                    resolution.review_status,
                    effective_action,
                ),
            )
            apply_row_fill(
                ws,
                new_row,
                action_fill(effective_action),
                excluded_columns=(columns["function"],),
            )
            add_comment(ws.cell(new_row, columns["subcategory"]), action.get("rationale") or "")
            record_merged_row(stats, effective_action, new_row)
            if resolution.requires_review:
                add_merge_review_event(
                    review_events,
                    change_id,
                    "Prototipo",
                    action_name,
                    effective_action,
                    new_row,
                    subcategory_id,
                    resolution.review_status,
                )
            if subcategory_id:
                row_index.setdefault(subcategory_id, new_row)

    return finish_workbook_merge(wb, ws, data, stats, review_events, output_path)


def merge_controls(
    controls_path: Path,
    output_path: Path,
    data: dict[str, Any],
    controls_filter_column: str | None,
    controls_prototype_flag_column: str | None,
) -> dict[str, int]:
    wb = load_workbook(controls_path)
    ws = wb[wb.sheetnames[0]]
    columns = find_controls_columns(
        ws,
        controls_filter_column,
        controls_prototype_flag_column,
    )
    header_row = int(columns["header_row"])
    main_table_last_col = ws.max_column
    ai_cols = ensure_columns(
        ws,
        [
            "AI Azione",
            "AI Change ID",
            "AI Punteggio confidenza",
            "AI Controllo proposto",
            "AI Subcategories proposte",
            "AI Riferimenti proposti",
            "AI Priorità",
            "AI Delegato a",
            "AI Razionale",
            "AI Stato revisione",
        ],
        header_row=header_row,
        separator_before=True,
        extend_body_format=False,
        borderless=True,
    )
    possible_separator_col = min(ai_cols.values()) - 1
    controls_separator_col = (
        possible_separator_col
        if is_blank_column(ws, possible_separator_col, header_row)
        else None
    )
    control_row_index = index_control_rows(
        ws,
        int(columns["subcategories"]),
        header_row,
    )
    controls_template_row = find_last_data_row(ws, int(columns["control"]), header_row)
    controls_next_row = find_last_value_row(ws, header_row, main_table_last_col) + 1
    stats = new_merge_stats()
    review_events: list[MergeReviewEvent] = []
    row_actions: dict[int, str] = {}

    for change in iter_change_contexts(data):
        change_id = value_to_text(change.get("change_id"))
        for action in change.get("control_actions", []):
            action_name = value_to_text(action.get("action"))
            if action_name == "no_change":
                continue

            resolution = RowResolution(None, "Nuovo controllo proposto")
            if action_name != "new":
                resolution = resolve_controls_target_row(
                    ws,
                    int(columns["subcategories"]),
                    control_row_index,
                    action.get("existing_control_row"),
                    action.get("subcategories_with_weight"),
            )
            effective_action = confidence_aware_merge_action(
                action,
                effective_merge_action(action_name, resolution.requires_review),
            )
            subject = (
                action.get("existing_control_id")
                or action.get("subcategories_with_weight")
                or action.get("control_text")
            )

            if resolution.row_idx and 1 <= resolution.row_idx <= ws.max_row:
                row_idx = resolution.row_idx
                effective_action = confidence_aware_merge_action(
                    action,
                    effective_row_action(
                        row_actions,
                        row_idx,
                        action_name,
                        resolution.requires_review,
                    ),
                )
                append_ai_values(
                    ws,
                    row_idx,
                    ai_cols,
                    control_ai_values(
                        action,
                        change_id,
                        resolution.review_status,
                        effective_action,
                    ),
                )
                apply_row_fill(ws, row_idx, action_fill(effective_action))
                add_comment(ws.cell(row_idx, int(columns["control"])), action.get("rationale") or "")
                record_merged_row(stats, effective_action, row_idx)
                if resolution.requires_review:
                    add_merge_review_event(
                        review_events,
                        change_id,
                        "Controlli",
                        action_name,
                        effective_action,
                        row_idx,
                        subject,
                        resolution.review_status,
                    )
                continue

            if action_name not in APPENDABLE_ACTIONS:
                if action_name == "deprecate":
                    add_merge_review_event(
                        review_events,
                        change_id,
                        "Controlli",
                        action_name,
                        effective_action,
                        None,
                        subject,
                        f"Deprecazione non applicata: {resolution.review_status}",
                    )
                continue

            new_row = controls_next_row
            controls_next_row += 1
            effective_action = confidence_aware_merge_action(
                action,
                effective_row_action(
                    row_actions,
                    new_row,
                    action_name,
                    resolution.requires_review,
                ),
            )
            copy_row_format(ws, controls_template_row, new_row)
            id_col = int(columns["id"] or 1)
            ws.cell(new_row, id_col).value = f"=ROW(A{new_row - 1})"
            ws.cell(new_row, int(columns["control"])).value = action.get("control_text")
            ws.cell(new_row, int(columns["subcategories"])).value = action.get("subcategories_with_weight")
            target_reference = columns.get("target_reference")
            if target_reference:
                ws.cell(new_row, int(target_reference)).value = action.get("references")
            target_flag = columns.get("target_flag")
            if target_flag:
                ws.cell(new_row, int(target_flag)).value = "X "
            append_ai_values(
                ws,
                new_row,
                ai_cols,
                control_ai_values(
                    action,
                    change_id,
                    resolution.review_status,
                    effective_action,
                ),
            )
            apply_row_fill(ws, new_row, action_fill(effective_action))
            add_comment(ws.cell(new_row, int(columns["control"])), action.get("rationale") or "")
            record_merged_row(stats, effective_action, new_row)
            if resolution.requires_review:
                add_merge_review_event(
                    review_events,
                    change_id,
                    "Controlli",
                    action_name,
                    effective_action,
                    new_row,
                    subject,
                    resolution.review_status,
                )

    if controls_separator_col:
        format_separator_column(ws, controls_separator_col, header_row)
    clear_column_borders(ws, ai_cols.values(), header_row)
    return finish_workbook_merge(wb, ws, data, stats, review_events, output_path)


def positive_path(value: str | Path | None) -> Path | None:
    if not value:
        return None
    return Path(value)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Update a user-provided contextualization prototype and controls workbook "
            "from annotated regulatory changes."
        )
    )
    parser.add_argument("--annotations-json", required=True, type=Path)
    parser.add_argument(
        "--framework-core",
        required=True,
        type=Path,
        help="Framework core workbook matching the supplied prototype and controls.",
    )
    parser.add_argument(
        "--prototype-xlsx",
        type=positive_path,
        required=True,
        help="Contextualization prototype workbook to update.",
    )
    parser.add_argument(
        "--prototype-reference-column",
        help=(
            "Prototype column containing references to the selected source. "
            "Default: autodetect the Circular 285 reference column for compatibility."
        ),
    )
    parser.add_argument(
        "--controls-xlsx",
        type=positive_path,
        required=True,
        help="Controls workbook to update.",
    )
    parser.add_argument("--prompt", default=DEFAULT_PROMPT, type=Path)
    parser.add_argument("--query-script", default=DEFAULT_QUERY_SCRIPT, type=Path)
    parser.add_argument("--llm", choices=["gemini", "openai"], required=False)
    parser.add_argument(
        "--output-language",
        default=DEFAULT_OUTPUT_LANGUAGE,
        help=(
            "Language for generated workbook-facing content and the LLM response. "
            f"Default: {DEFAULT_OUTPUT_LANGUAGE} (Italian)."
        ),
    )
    parser.add_argument("--output-dir", default=Path("contextualization_output"), type=Path)
    parser.add_argument("--output-json", type=Path)
    parser.add_argument(
        "--contextualization-json",
        type=Path,
        help="Use an already generated contextualization JSON and skip the LLM call.",
    )
    parser.add_argument(
        "--controls-reference-column",
        "--controls-filter-column",
        dest="controls_filter_column",
        help=(
            "Controls column containing references for the selected prototype, for example "
            "D for Payment Services or G for the General Prototype. Default: autodetect "
            "the General Prototype references column."
        ),
    )
    parser.add_argument(
        "--controls-prototype-flag-column",
        help=(
            "Controls column marking membership in the selected prototype, for example H "
            "for Payment Services or K for the General Prototype. Default: autodetect the "
            "General Prototype flag column."
        ),
    )
    parser.add_argument(
        "--controls-reference-pattern",
        help=(
            "Case-insensitive regular expression that selected control references must match, "
            "for example '2015\\s*[/.-]\\s*2366' for PSD2."
        ),
    )
    parser.add_argument(
        "--no-circ285-filter",
        action="store_true",
        help=(
            "Disable the legacy default Circular 285 reference filter. Retained for backward "
            "compatibility; prefer --controls-reference-pattern for non-Circular sources."
        ),
    )
    parser.add_argument(
        "--max-controls-context",
        type=int,
        default=180,
        help="Maximum number of existing controls included in the LLM context.",
    )
    parser.add_argument(
        "--prepare-only",
        action="store_true",
        help="Prepare the bundle and schema without calling the model or generating workbooks.",
    )
    parser.add_argument(
        "--no-merge",
        action="store_true",
        help="Generate/use the contextualization JSON but do not create Excel workbooks.",
    )
    parser.add_argument(
        "--clean",
        action="store_true",
        help=(
            "Also create publication-ready prototype and controls workbooks with no "
            "AI notes, generated highlights, or AI change-log sheet. Low-confidence and "
            "review-needed actions are not applied."
        ),
    )
    parser.add_argument(
        "--confidence-self-consistency",
        action="store_true",
        help=(
            "Run a second independent contextualization query and add up to 10 confidence "
            "points when it proposes the same action and FNCDP subcategories."
        ),
    )
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> None:
    required = [args.annotations_json, args.framework_core]
    if not args.prepare_only and (
        not args.contextualization_json or args.confidence_self_consistency
    ):
        required.extend([args.prompt, args.query_script])
    for path in required:
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

    if args.contextualization_json and not args.contextualization_json.exists():
        raise FileNotFoundError(f"Contextualization JSON not found: {args.contextualization_json}")

    if args.prototype_xlsx is None:
        raise ValueError("--prototype-xlsx is required.")
    if args.controls_xlsx is None:
        raise ValueError("--controls-xlsx is required.")
    for workbook_path in [args.prototype_xlsx, args.controls_xlsx]:
        if not workbook_path.exists():
            raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    if args.controls_reference_pattern and args.no_circ285_filter:
        raise ValueError(
            "--controls-reference-pattern cannot be combined with --no-circ285-filter."
        )
    if args.controls_reference_pattern:
        try:
            re.compile(args.controls_reference_pattern, re.IGNORECASE)
        except re.error as error:
            raise ValueError(f"Invalid --controls-reference-pattern: {error}") from error

    for column_name, column_ref in [
        ("--prototype-reference-column", args.prototype_reference_column),
        ("--controls-reference-column", args.controls_filter_column),
        ("--controls-prototype-flag-column", args.controls_prototype_flag_column),
    ]:
        try:
            parse_column_ref(column_ref)
        except ValueError as error:
            raise ValueError(f"Invalid {column_name}: {column_ref}") from error

    if not value_to_text(args.output_language):
        raise ValueError("--output-language cannot be empty.")

    if args.prepare_only and args.confidence_self_consistency:
        raise ValueError("--confidence-self-consistency cannot be used with --prepare-only.")


def main() -> None:
    args = parse_args()
    try:
        validate_args(args)
    except (FileNotFoundError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        raise SystemExit(1) from error
    args.output_dir.mkdir(parents=True, exist_ok=True)

    bundle, framework_lookup = build_input_bundle(args)
    try:
        allowed_subcategory_ids = collect_framework_subcategory_ids(framework_lookup)
        allowed_control_subcategory_ids = sorted(
            set(allowed_subcategory_ids) | set(collect_workbook_subcategory_ids(bundle))
        )
        expected_change_ids = collect_expected_change_ids(bundle)
    except ValueError as error:
        print(f"Error: {error}", file=sys.stderr)
        raise SystemExit(1) from error

    bundle_path = args.output_dir / "contextualization_input_bundle.json"
    schema_path = args.output_dir / "update_contextualization_schema.generated.json"
    scored_schema_path = args.output_dir / "update_contextualization_scored_schema.generated.json"
    output_json = args.output_json or (args.output_dir / "contextualization_updates.json")
    output_schema = schema(
        allowed_subcategory_ids,
        expected_change_ids,
        args.output_language,
    )
    scored_output_schema = schema(
        allowed_subcategory_ids,
        expected_change_ids,
        args.output_language,
        include_confidence_fields=True,
    )

    write_json(bundle_path, bundle)
    write_json(schema_path, output_schema)
    write_json(scored_schema_path, scored_output_schema)
    print(f"Created input bundle: {bundle_path}")
    print(f"Created model schema: {schema_path}")
    print(f"Created scored-output schema: {scored_schema_path}")

    if args.prepare_only:
        print("Prepare-only completed. No LLM call was made.")
        return

    if args.contextualization_json:
        output_json = args.contextualization_json
    else:
        run_query(args, bundle_path, schema_path, output_json)

    data = read_json(output_json)
    normalize_contextualization_subcategory_ids(data)
    normalize_legacy_source_reference_fields(data)
    if args.contextualization_json:
        remove_confidence_fields(data)
    validate_contextualization(
        data,
        output_schema,
        allowed_subcategory_ids,
        expected_change_ids,
        allowed_control_subcategory_ids,
    )
    print(f"Contextualization JSON: {output_json}")

    self_consistency = run_self_consistency_check(
        args,
        bundle_path,
        schema_path,
        output_schema,
        allowed_subcategory_ids,
        expected_change_ids,
        allowed_control_subcategory_ids,
    )
    annotate_confidence_assessments(
        data,
        args.prototype_xlsx,
        args.controls_xlsx,
        framework_lookup,
        args.controls_filter_column,
        args.controls_prototype_flag_column,
        args.prototype_reference_column,
        self_consistency,
    )
    validate_contextualization(
        data,
        scored_output_schema,
        allowed_subcategory_ids,
        expected_change_ids,
        allowed_control_subcategory_ids,
    )
    scored_output_json = (
        args.output_dir / f"{output_json.stem}_scored.json"
        if args.contextualization_json
        else output_json
    )
    write_json(scored_output_json, data)
    print(f"Scored contextualization JSON: {scored_output_json}")

    if args.no_merge:
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    prototype_output = args.output_dir / f"{args.prototype_xlsx.stem}_{timestamp}.xlsx"
    prototype_stats = merge_prototype(
        args.prototype_xlsx,
        prototype_output,
        data,
        framework_lookup,
        args.prototype_reference_column,
    )
    print(f"Created prototype workbook: {prototype_output}")
    print(
        "Prototype merge statistics: "
        f"new={prototype_stats['new_rows']}, "
        f"modified={prototype_stats['modified_rows']}, "
        f"deprecated={prototype_stats['deprecated_rows']}, "
        f"review_needed={prototype_stats['review_needed_rows']}"
    )

    controls_output = args.output_dir / f"{args.controls_xlsx.stem}_{timestamp}.xlsx"
    controls_stats = merge_controls(
        args.controls_xlsx,
        controls_output,
        data,
        args.controls_filter_column,
        args.controls_prototype_flag_column,
    )
    print(f"Created controls workbook: {controls_output}")
    print(
        "Controls merge statistics: "
        f"new={controls_stats['new_rows']}, "
        f"modified={controls_stats['modified_rows']}, "
        f"deprecated={controls_stats['deprecated_rows']}, "
        f"review_needed={controls_stats['review_needed_rows']}"
    )

    if args.clean:
        clean_prototype_output = args.output_dir / (
            f"{args.prototype_xlsx.stem}_clean_{timestamp}.xlsx"
        )
        clean_prototype_stats = merge_clean_prototype(
            args.prototype_xlsx,
            clean_prototype_output,
            data,
            framework_lookup,
            args.prototype_reference_column,
        )
        print(f"Created clean prototype workbook: {clean_prototype_output}")
        print(
            "Clean prototype merge statistics: "
            f"new={clean_prototype_stats['new_rows']}, "
            f"modified={clean_prototype_stats['modified_rows']}"
        )

        clean_controls_output = args.output_dir / (
            f"{args.controls_xlsx.stem}_clean_{timestamp}.xlsx"
        )
        clean_controls_stats = merge_clean_controls(
            args.controls_xlsx,
            clean_controls_output,
            data,
            args.controls_filter_column,
            args.controls_prototype_flag_column,
        )
        print(f"Created clean controls workbook: {clean_controls_output}")
        print(
            "Clean controls merge statistics: "
            f"new={clean_controls_stats['new_rows']}, "
            f"modified={clean_controls_stats['modified_rows']}"
        )


if __name__ == "__main__":
    main()
